#!/usr/bin/env python3
from __future__ import annotations

import argparse
import sys
from pathlib import Path

try:
    from docx import Document
except ImportError as exc:  # pragma: no cover
    raise SystemExit('python-docx is required. Install requirements-dev.txt') from exc


def add_markdown_line(document: Document, line: str) -> None:
    stripped = line.strip()
    if not stripped:
        document.add_paragraph('')
        return
    if stripped.startswith('# '):
        document.add_heading(stripped[2:].strip(), level=1)
        return
    if stripped.startswith('## '):
        document.add_heading(stripped[3:].strip(), level=2)
        return
    if stripped.startswith('### '):
        document.add_heading(stripped[4:].strip(), level=3)
        return
    if stripped in {'Reviewer Reports on the Initial Version', "Referees' comments:"}:
        document.add_heading(stripped, level=1 if stripped.startswith('Reviewer') else 2)
        return
    if stripped.startswith('Referee #'):
        document.add_heading(stripped, level=1)
        return
    if stripped in {'Major comments:', 'Specific comments:', 'Overall assessment:'}:
        document.add_heading(stripped, level=2)
        return
    document.add_paragraph(stripped)


def render_docx(input_md: Path, output_docx: Path) -> None:
    if not input_md.is_file():
        raise FileNotFoundError(f'input markdown not found: {input_md}')
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for line in input_md.read_text(encoding='utf-8').splitlines():
        add_markdown_line(document, line)
    document.save(output_docx)


def main() -> int:
    parser = argparse.ArgumentParser(description='Render a Nature-style review markdown file to DOCX.')
    parser.add_argument('input_md', help='input Markdown review report')
    parser.add_argument('output_docx', help='output DOCX path')
    args = parser.parse_args()
    render_docx(Path(args.input_md), Path(args.output_docx))
    print(f'wrote {args.output_docx}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
