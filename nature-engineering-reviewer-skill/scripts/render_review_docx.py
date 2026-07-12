from __future__ import annotations

import argparse
from pathlib import Path

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    try:
        from docx import Document
    except ImportError as exc:
        raise SystemExit("Install python-docx==1.1.2 to render DOCX") from exc
    text = args.input.read_text(encoding="utf-8")
    document = Document()
    for block in text.split("\n\n"):
        stripped = block.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
        else:
            document.add_paragraph(stripped)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(args.output)
