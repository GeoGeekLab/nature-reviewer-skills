from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


def add_markdown_line(doc: Document, line: str) -> None:
    if line.startswith('# '):
        doc.add_heading(line[2:].strip(), level=1)
    elif line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=2)
    elif line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=3)
    elif line.startswith('- '):
        doc.add_paragraph(line[2:].strip(), style='List Bullet')
    elif line.strip().startswith(tuple(f'{i}. ' for i in range(1, 10))):
        numbered = line.strip().split('. ', 1)[1] if '. ' in line.strip() else line.strip()
        doc.add_paragraph(numbered, style='List Number')
    elif line.strip():
        doc.add_paragraph(line.strip())


def render_docx(markdown_path: Path, output_path: Path) -> None:
    if not markdown_path.exists():
        raise FileNotFoundError(f'Markdown file not found: {markdown_path}')
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal'].font.size = Pt(10.5)
    for line in markdown_path.read_text(encoding='utf-8').splitlines():
        add_markdown_line(doc, line)
    doc.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description='Render a Markdown review report to DOCX.')
    parser.add_argument('markdown_path', type=Path)
    parser.add_argument('output_path', type=Path)
    args = parser.parse_args()
    render_docx(args.markdown_path, args.output_path)
    print(f'wrote {args.output_path}')
    return 0

if __name__ == '__main__':
    sys.exit(main())
