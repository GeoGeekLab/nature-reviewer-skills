"""Render a Nature-style referee Markdown report to Word DOCX."""

from __future__ import annotations

import argparse
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
except ImportError as exc:  # pragma: no cover - exercised only without dependency.
    raise SystemExit(
        "Missing dependency `python-docx`. Install it with: "
        "python -m pip install python-docx"
    ) from exc

REPORT_BOLD_LINES: frozenset[str] = frozenset(
    {
        "Reviewer Reports on the Initial Version:",
        "Referees' comments:",
        "Decision tendency:",
    }
)


def render_markdown_to_docx(markdown_path: Path, docx_path: Path) -> None:
    """Render a conservative Markdown referee report to DOCX.

    The renderer intentionally supports a small Markdown subset used by the skill's
    peer-review-file template: headings, bullet lists, numbered comments, and plain
    paragraphs. This keeps the Word output predictable for GitHub use and avoids
    hidden formatting dependencies.
    """
    markdown_text = _read_markdown(markdown_path)
    document = _create_document()

    for line in markdown_text.splitlines():
        add_markdown_line(document, line)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(docx_path)


def add_markdown_line(document: Document, line: str) -> None:
    """Add one supported Markdown line to a Word document."""
    stripped = line.strip()
    if not stripped:
        document.add_paragraph()
        return

    if _add_heading(document, stripped):
        return
    if stripped.startswith("- "):
        document.add_paragraph(stripped[2:].strip(), style="List Bullet")
        return
    if _is_numbered_comment(stripped):
        document.add_paragraph(stripped, style="List Number")
        return
    if _is_report_heading(stripped):
        _add_bold_paragraph(document, stripped)
        return

    document.add_paragraph(stripped)


def _read_markdown(markdown_path: Path) -> str:
    """Read and validate the Markdown source file."""
    if not markdown_path.exists():
        raise FileNotFoundError(f"Input Markdown file does not exist: {markdown_path}")
    if not markdown_path.is_file():
        raise ValueError(f"Input path is not a file: {markdown_path}")
    return markdown_path.read_text(encoding="utf-8")


def _create_document() -> Document:
    """Create a Word document with conservative manuscript-review styling."""
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    for heading_name in ("Heading 1", "Heading 2", "Heading 3"):
        heading_style = document.styles[heading_name]
        heading_style.font.name = "Times New Roman"

    return document


def _add_heading(document: Document, text: str) -> bool:
    """Add a Markdown heading when text starts with a supported heading marker."""
    heading_levels = (("### ", 3), ("## ", 2), ("# ", 1))
    for prefix, level in heading_levels:
        if text.startswith(prefix):
            paragraph = document.add_heading(text[len(prefix) :].strip(), level=level)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            return True
    return False


def _is_report_heading(text: str) -> bool:
    """Return whether text is a Nature-style referee-report heading."""
    return text in REPORT_BOLD_LINES or (
        text.startswith("Referee #") and text.endswith(":")
    )


def _add_bold_paragraph(document: Document, text: str) -> None:
    """Add a bold plain-text paragraph."""
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True


def _is_numbered_comment(text: str) -> bool:
    """Return whether text begins with a simple numbered comment marker."""
    marker, separator, _rest = text.partition(".")
    return separator == "." and marker.isdigit()


def parse_args() -> argparse.Namespace:
    """Parse command-line arguments."""
    parser = argparse.ArgumentParser(
        description="Render a Nature-style referee Markdown report to DOCX."
    )
    parser.add_argument("input_md", help="Input Markdown report.")
    parser.add_argument("output_docx", help="Output DOCX path.")
    return parser.parse_args()


def main() -> int:
    """Render a Markdown report to DOCX and return a shell status code."""
    args = parse_args()
    render_markdown_to_docx(Path(args.input_md), Path(args.output_docx))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
