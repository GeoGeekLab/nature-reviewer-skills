from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.shared import Pt


def add_markdown_line(doc: Document, line: str) -> None:
    if line.startswith('# '):
        doc.add_heading(line[2:].strip(), level=1)
    elif line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=2)
    elif line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=3)
    elif line.startswith('- '):
        doc.add_paragraph(line[2:].strip(), style='List Bullet')
    elif line.strip() == '':
        doc.add_paragraph('')
    else:
        doc.add_paragraph(line)


def render(markdown_path: Path, docx_path: Path) -> None:
    text = markdown_path.read_text(encoding='utf-8')
    doc = Document()
    styles = doc.styles
    styles['Normal'].font.name = 'Times New Roman'
    styles['Normal'].font.size = Pt(11)
    for line in text.splitlines():
        add_markdown_line(doc, line)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


def main() -> None:
    parser = argparse.ArgumentParser(description='Render a Markdown review report to DOCX.')
    parser.add_argument('markdown', type=Path)
    parser.add_argument('docx', type=Path)
    args = parser.parse_args()
    if not args.markdown.is_file():
        raise SystemExit(f'Markdown input not found: {args.markdown}')
    render(args.markdown, args.docx)
    print(f'Wrote {args.docx}')


if __name__ == '__main__':
    main()
