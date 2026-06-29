from __future__ import annotations

import argparse
from pathlib import Path

try:
    from docx import Document
except ImportError as exc:  # pragma: no cover
    raise SystemExit('python-docx is required: pip install python-docx') from exc


def render(markdown_path: Path, output_path: Path) -> None:
    text = markdown_path.read_text(encoding='utf-8')
    doc = Document()
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith('# '):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
        else:
            doc.add_paragraph(stripped)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description='Render a Markdown review report to DOCX.')
    parser.add_argument('markdown', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    if not args.markdown.exists():
        raise SystemExit(f'Markdown file not found: {args.markdown}')
    render(args.markdown, args.output)
    print(f'Wrote {args.output}')

if __name__ == '__main__':
    main()
